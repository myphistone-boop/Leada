"""
=============================================================================
GENERATEUR DE LEADS GOOGLE MAPS
=============================================================================
Recherche des entreprises sur Google Maps et exporte les données en Excel.

Usage:
    pip install -r requirements.txt
    playwright install chromium
    python lead_generator.py
=============================================================================
"""

import re
from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from datetime import datetime
from urllib.parse import quote

# =============================================================================
# CONFIGURATION
# =============================================================================

DEBUG = False  # Mettre True pour voir le navigateur
SCROLL_PAUSE = 2  # Pause entre chaque scroll (secondes)
MAX_SCROLLS = 20  # Nombre max de scrolls


# =============================================================================
# SCRAPING GOOGLE MAPS
# =============================================================================

def rechercher_entreprises(activite: str, localisation: str, max_resultats: int = 100, sans_site_uniquement: bool = False) -> list:
    """
    Scrape Google Maps pour trouver des entreprises.

    Args:
        activite: Type d'entreprise (ex: "plombier")
        localisation: Zone géographique (ex: "Lyon")
        max_resultats: Nombre max de résultats
        sans_site_uniquement: Si True, garde uniquement les entreprises sans site web

    Returns:
        Liste des entreprises avec leurs infos
    """
    print(f"\n🔍 Recherche: '{activite}' à '{localisation}'")
    if sans_site_uniquement:
        print(f"🎯 Filtre: entreprises SANS site web uniquement")

    # Construire l'URL de recherche
    query = quote(f"{activite} {localisation}")
    url = f"https://www.google.com/maps/search/{query}"

    entreprises = []

    with sync_playwright() as p:
        # Lancer le navigateur
        browser = p.chromium.launch(headless=not DEBUG)
        page = browser.new_page()

        print(f"🌐 Chargement de Google Maps...")
        page.goto(url, timeout=60000)

        # Accepter les cookies RGPD
        cookies_buttons = ["Tout accepter", "Accept all", "Accepter tout", "J'accepte", "Agree"]
        for btn_text in cookies_buttons:
            try:
                page.click(f"button:has-text('{btn_text}')", timeout=2000)
                print(f"✓ Cookies acceptés")
                page.wait_for_timeout(1000)
                break
            except:
                continue

        # Attendre les résultats
        print(f"⏳ Attente du chargement...")
        try:
            page.wait_for_selector("div[role='feed']", timeout=30000)
        except:
            print(f"❌ Impossible de charger les résultats")
            page.screenshot(path="debug_screenshot.png")
            print(f"📸 Screenshot sauvé: debug_screenshot.png")
            browser.close()
            return []

        print(f"📜 Scroll pour charger plus de résultats...")

        # Scroller pour charger plus
        feed = page.query_selector("div[role='feed']")
        previous_count = 0

        for i in range(MAX_SCROLLS):
            feed.evaluate("el => el.scrollTop = el.scrollHeight")
            page.wait_for_timeout(SCROLL_PAUSE * 1000)

            items = page.query_selector_all("div[role='feed'] > div > div[jsaction]")
            current_count = len(items)

            print(f"  Scroll {i+1}/{MAX_SCROLLS} - {current_count} résultats")

            if current_count >= max_resultats * 2 or current_count == previous_count:
                break
            previous_count = current_count

        # Extraire les données
        print(f"\n📊 Extraction des données...")
        items = page.query_selector_all("div[role='feed'] > div > div[jsaction]")

        for item in items:
            if len(entreprises) >= max_resultats:
                break

            try:
                item.click()
                page.wait_for_timeout(2000)  # Attente pour charger les détails

                entreprise = extraire_details(page)

                # Validation: nom ET téléphone obligatoires
                if not entreprise.get("nom") or not entreprise.get("telephone"):
                    nom = entreprise.get("nom", "???")
                    print(f"  ⚠ {nom} - données incomplètes (pas de tél)")
                    continue

                # Filtre sans site web si demandé
                if sans_site_uniquement and entreprise.get("site_web"):
                    print(f"  ✗ {entreprise['nom']} (a un site web)")
                    continue

                entreprises.append(entreprise)
                status = "sans site" if not entreprise.get("site_web") else "avec site"
                print(f"  ✓ {entreprise['nom']} | {entreprise['telephone']} ({status})")

            except:
                continue

        browser.close()

    print(f"\n✅ {len(entreprises)} entreprises extraites")
    return entreprises


def extraire_details(page) -> dict:
    """Extrait les détails d'une entreprise depuis le panneau Google Maps."""
    entreprise = {
        "nom": "",
        "adresse": "",
        "telephone": "",
        "site_web": "",
        "note": "",
        "nb_avis": "",
        "categorie": ""
    }

    try:
        # Attendre que le panneau de détails charge
        page.wait_for_selector("h1", timeout=5000)

        # === NOM (plusieurs sélecteurs) ===
        for selector in ["h1.DUwDvf", "h1[class*='header']", "h1"]:
            nom_el = page.query_selector(selector)
            if nom_el:
                nom = nom_el.inner_text().strip()
                if nom and len(nom) > 1:
                    entreprise["nom"] = nom
                    break

        # === NOTE ET AVIS ===
        note_el = page.query_selector("div.F7nice span[aria-hidden='true']")
        if note_el:
            entreprise["note"] = note_el.inner_text().strip()

        # Chercher le nombre d'avis avec plusieurs patterns
        avis_selectors = [
            "div.F7nice span[aria-label*='avis']",
            "span[aria-label*='review']",
            "span[aria-label*='avis']"
        ]
        for sel in avis_selectors:
            avis_el = page.query_selector(sel)
            if avis_el:
                avis_text = avis_el.get_attribute("aria-label") or ""
                nb = re.search(r"(\d[\d\s]*)", avis_text.replace("\u202f", "").replace(" ", ""))
                if nb:
                    entreprise["nb_avis"] = nb.group(1).replace(" ", "")
                    break

        # === CATEGORIE ===
        cat_el = page.query_selector("button[jsaction*='category']")
        if cat_el:
            entreprise["categorie"] = cat_el.inner_text().strip()

        # === ADRESSE, TELEPHONE, SITE WEB via les boutons ===
        # Méthode 1: boutons avec data-item-id
        buttons = page.query_selector_all("button[data-item-id]")
        for btn in buttons:
            item_id = btn.get_attribute("data-item-id") or ""
            aria = btn.get_attribute("aria-label") or ""

            if "address" in item_id and not entreprise["adresse"]:
                entreprise["adresse"] = aria.replace("Adresse:", "").replace("Address:", "").strip()

            elif "phone" in item_id and not entreprise["telephone"]:
                # Extraire juste le numéro
                tel = aria.replace("Téléphone:", "").replace("Phone:", "").strip()
                entreprise["telephone"] = tel

        # Méthode 2: chercher le téléphone dans les liens/texte si pas trouvé
        if not entreprise["telephone"]:
            # Chercher un lien tel:
            tel_link = page.query_selector("a[href^='tel:']")
            if tel_link:
                href = tel_link.get_attribute("href") or ""
                entreprise["telephone"] = href.replace("tel:", "").strip()

            # Ou chercher dans le texte avec regex
            if not entreprise["telephone"]:
                page_text = page.inner_text()
                tel_match = re.search(r"(\+?\d{1,3}[\s.-]?\(?\d{2,4}\)?[\s.-]?\d{2,4}[\s.-]?\d{2,4}[\s.-]?\d{0,4})", page_text)
                if tel_match:
                    entreprise["telephone"] = tel_match.group(1).strip()

        # === SITE WEB ===
        site_selectors = [
            "a[data-item-id='authority']",
            "a[aria-label*='site']",
            "a[aria-label*='Site']",
            "a[aria-label*='website']"
        ]
        for sel in site_selectors:
            site_el = page.query_selector(sel)
            if site_el:
                href = site_el.get_attribute("href")
                if href and "google" not in href:
                    entreprise["site_web"] = href
                    break

    except Exception as e:
        pass

    return entreprise


# =============================================================================
# EXPORT EXCEL
# =============================================================================

def exporter_excel(entreprises: list, nom_fichier: str = None) -> str:
    """Exporte les leads dans un fichier Excel."""
    if not nom_fichier:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        nom_fichier = f"leads_{timestamp}.xlsx"

    print(f"\n📊 Export: {nom_fichier}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    # En-têtes
    headers = ["Nom", "Catégorie", "Adresse", "Téléphone", "Site Web", "Note", "Avis"]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    # Données
    for row, e in enumerate(entreprises, 2):
        ws.cell(row=row, column=1, value=e.get("nom", ""))
        ws.cell(row=row, column=2, value=e.get("categorie", ""))
        ws.cell(row=row, column=3, value=e.get("adresse", ""))
        ws.cell(row=row, column=4, value=e.get("telephone", ""))
        ws.cell(row=row, column=5, value=e.get("site_web", ""))
        ws.cell(row=row, column=6, value=e.get("note", ""))
        ws.cell(row=row, column=7, value=e.get("nb_avis", ""))

    # Largeur colonnes
    for col, width in enumerate([30, 20, 45, 18, 40, 8, 10], 1):
        ws.column_dimensions[chr(64 + col)].width = width

    wb.save(nom_fichier)
    print(f"✅ Fichier créé: {nom_fichier}")
    return nom_fichier


# =============================================================================
# FONCTION PRINCIPALE
# =============================================================================

def generer_leads(activite: str, localisation: str, max_resultats: int = 100, sans_site_uniquement: bool = False) -> str:
    """Génère des leads et les exporte en Excel."""
    print("=" * 60)
    print("GENERATEUR DE LEADS GOOGLE MAPS")
    print("=" * 60)

    # Scraper Google Maps
    entreprises = rechercher_entreprises(activite, localisation, max_resultats, sans_site_uniquement)

    if not entreprises:
        print("❌ Aucune entreprise trouvée")
        return None

    # Exporter
    fichier = exporter_excel(entreprises)

    print("\n" + "=" * 60)
    print("TERMINÉ!")
    print("=" * 60)

    return fichier


# =============================================================================
# POINT D'ENTREE
# =============================================================================

if __name__ == "__main__":
    print("\n" + "=" * 60)
    print("CONFIGURATION")
    print("=" * 60)

    activite = input("\n📌 Activité (ex: plombier, restaurant): ").strip()
    localisation = input("📍 Localisation (ex: Lyon, Paris 15): ").strip()

    # Option sans site web
    sans_site = input("🚫 Uniquement les entreprises SANS site web ? (o/N): ").strip().lower()
    sans_site_uniquement = sans_site in ["o", "oui", "y", "yes"]

    max_res = input("📊 Nombre max de résultats (défaut: 100): ").strip()
    max_resultats = int(max_res) if max_res.isdigit() else 100

    generer_leads(activite, localisation, max_resultats, sans_site_uniquement)
